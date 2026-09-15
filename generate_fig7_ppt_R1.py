"""
Fig. 7 (revised, R1): Slashdot subgraph topology on the 140-node strongly
connected core, drawn in exactly the style of the original hand-finished
PowerPoint figure (Fig7.pptx, slide 2).

The original slide is used as a template: node positions, node size, white
outline and shadow, edge widths/transparency/dash pattern and the bottom
legend are all kept.  The 9 nodes that are not in the strongly connected core
of the pruned graph are removed, the edges are redrawn from the 140-node
adjacency (same subsampling rule as the original: at most 300 positive edges,
seed 42), and the node counts in the legend are updated.

Output: Fig7_R1.pptx (editable).  Export slide 1 from PowerPoint as PDF/PNG.
"""
import os, sys, copy
import numpy as np
import scipy.sparse as sp
from pptx import Presentation
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
import RVGB_slashdot_experiment as SE
from revision_analysis import scc_info

TEMPLATE = os.path.join(HERE, 'Fig7.pptx')   # original hand-finished figure
OUTPUT = os.path.join(HERE, 'Fig7_R1.pptx')
NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
NODE_EXT = 146304          # EMU, diameter of a node oval in the original slide
MAX_DISPLAY = 300          # same edge subsampling as the original figure


def q(tag):
    p, t = tag.split(':'); return '{%s}%s' % (NS[p], t)


def spectral_layout(A):
    n = A.shape[0]
    A_sym = (np.abs(A) + np.abs(A.T)) / 2
    D = A_sym.sum(axis=1); D[D == 0] = 1
    L = np.eye(n) - np.diag(1 / np.sqrt(D)) @ A_sym @ np.diag(1 / np.sqrt(D))
    _, vec = np.linalg.eigh(L)
    return vec[:, 1], vec[:, 2]


def main():
    # ---- data: identical pipeline to experiment6_R1.py ----
    SE.download_dataset(); edges = SE.load_edge_list()
    A0, n0, _, _ = SE.extract_dense_subgraph(edges, target_size=150)
    comm0 = SE.sign_consistent_partition(A0)
    d0, ratio0, A1 = SE.assign_gauge_and_verify(A0, comm0, d_values=(1.0, -2.0))
    _, _, lab = scc_info(sp.csr_matrix(A1)); comp = np.argmax(np.bincount(lab))
    keep = np.where(lab == comp)[0]
    A2 = A1[np.ix_(keep, keep)]; d2 = d0[keep]; n2 = len(keep)
    V1 = [k for k in range(n2) if d2[k] > 0]; V2 = [k for k in range(n2) if d2[k] < 0]
    print(f'149-node pruned graph -> core: {n2} nodes, |V1|={len(V1)}, |V2|={len(V2)}')

    # ---- template slide ----
    prs = Presentation(TEMPLATE)
    sldIdLst = prs.slides._sldIdLst
    first = sldIdLst[0]; prs.part.drop_rel(first.rId); sldIdLst.remove(first)   # drop slide 1
    slide = prs.slides[0]
    grp = slide.shapes._spTree.find('.//p:grpSp', NS)
    sps = grp.findall('p:sp', NS); cxns = grp.findall('p:cxnSp', NS)
    ovals = [s for s in sps if s.find('.//a:prstGeom', NS).get('prst') == 'ellipse']
    # the generator added the 149 node ovals first and the two legend ovals last
    node_ovals, legend_ovals = ovals[:n0], ovals[n0:]
    assert len(legend_ovals) == 2 and all(int(o.find('.//a:off', NS).get('y')) > 5_000_000 for o in legend_ovals)
    texts = [s for s in sps if s.find('.//a:t', NS) is not None]

    # node k of the original graph -> oval (generator order: community 1 then 2)
    order = list(comm0[0]) + list(comm0[1])
    centre = {}
    for k, o in zip(order, node_ovals):
        off = o.find('.//a:off', NS); ext = o.find('.//a:ext', NS)
        centre[k] = (int(off.get('x')) + int(ext.get('cx')) // 2,
                     int(off.get('y')) + int(ext.get('cy')) // 2)
    # sanity check: the oval centres must be an affine image of the spectral layout
    px, py = spectral_layout(A1)
    X = np.array([[px[k], 1] for k in order]); Y = np.array([[py[k], 1] for k in order])
    cx = np.array([centre[k][0] for k in order]); cy = np.array([centre[k][1] for k in order])
    rx = np.linalg.lstsq(X, cx, rcond=None)[1]; ry = np.linalg.lstsq(Y, cy, rcond=None)[1]
    relx = np.sqrt(rx[0] / len(order)) / (cx.max() - cx.min()); rely = np.sqrt(ry[0] / len(order)) / (cy.max() - cy.min())
    print(f'layout check: relative RMS residual x={relx:.2e}, y={rely:.2e}')
    assert relx < 0.02 and rely < 0.02, 'oval order does not match the spectral layout'

    # ---- templates and clean-up ----
    t_oval = copy.deepcopy(node_ovals[0])
    t_black = copy.deepcopy(next(c for c in cxns if 'D62728' not in etree.tostring(c).decode()))
    t_red = copy.deepcopy(next(c for c in cxns if 'D62728' in etree.tostring(c).decode()))
    for el in node_ovals + cxns: grp.remove(el)
    insert_at = list(grp).index(legend_ovals[0])  # before the legend
    next_id = [max(int(e.get('id')) for e in slide.shapes._spTree.iter(q('p:cNvPr'))) + 1]

    def new_id(el, name):
        c = el.find('.//p:cNvPr', NS); c.set('id', str(next_id[0])); c.set('name', f'{name} {next_id[0]}'); next_id[0] += 1

    def add_cxn(tmpl, p0, p1, name):
        el = copy.deepcopy(tmpl); new_id(el, name)
        xfrm = el.find('.//a:xfrm', NS); off = xfrm.find('a:off', NS); ext = xfrm.find('a:ext', NS)
        (x0, y0), (x1, y1) = p0, p1
        off.set('x', str(min(x0, x1))); off.set('y', str(min(y0, y1)))
        ext.set('cx', str(abs(x1 - x0))); ext.set('cy', str(abs(y1 - y0)))
        for a in ('flipH', 'flipV'):
            if a in xfrm.attrib: del xfrm.attrib[a]
        if (x1 - x0) * (y1 - y0) < 0: xfrm.set('flipH', '1')
        return el

    def add_node(p, colour, name):
        el = copy.deepcopy(t_oval); new_id(el, name)
        off = el.find('.//a:off', NS); off.set('x', str(p[0] - NODE_EXT // 2)); off.set('y', str(p[1] - NODE_EXT // 2))
        el.find('./p:spPr/a:solidFill/a:srgbClr', NS).set('val', colour)
        return el

    # ---- edges of the 140-node core (same subsampling as the original) ----
    pos_e = [(i, j) for i in range(n2) for j in range(n2) if i != j and A2[i, j] > 0]
    neg_e = [(i, j) for i in range(n2) for j in range(n2) if i != j and A2[i, j] < 0]
    n_pos, n_neg = len(pos_e), len(neg_e)
    rng = np.random.RandomState(42)
    if len(pos_e) > MAX_DISPLAY: pos_e = [pos_e[k] for k in rng.choice(len(pos_e), MAX_DISPLAY, replace=False)]
    if len(neg_e) > MAX_DISPLAY: neg_e = [neg_e[k] for k in rng.choice(len(neg_e), MAX_DISPLAY, replace=False)]
    P = {i: centre[keep[i]] for i in range(n2)}
    new = [add_cxn(t_black, P[i], P[j], 'Connector') for i, j in pos_e]
    new += [add_cxn(t_red, P[i], P[j], 'Connector') for i, j in neg_e]
    new += [add_node(P[i], '1F77B4', 'Oval') for i in V1]
    new += [add_node(P[i], 'D62728', 'Oval') for i in V2]
    for k, el in enumerate(new): grp.insert(insert_at + k, el)

    # ---- legend text ----
    for t in texts:
        for r in t.iter(q('a:t')):
            if r.text.startswith('V₁'): r.text = f'V₁ (dᵢ=+1.0, {len(V1)} nodes)'
            if r.text.startswith('V₂'): r.text = f'V₂ (dᵢ=−2.0, {len(V2)} nodes)'
    prs.save(OUTPUT)
    print(f'{n2} nodes, {n_pos}+ / {n_neg}- edges ({len(pos_e)}+{len(neg_e)} drawn); saved {OUTPUT}')


if __name__ == '__main__':
    main()
